"""Round-trip a troubleshooter knowledge base between JSON and an editable Excel workbook.

  python excel_io.py export kb.json Content.xlsx images/      # JSON -> workbook (+ images written to folder)
  python excel_io.py import Content.xlsx images/ kb.json      # workbook (+ images folder) -> JSON
"""
import json, sys, os, base64, re
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="0E6F73")
INPUT_FILL = PatternFill("solid", fgColor="FFF9E6")

SHEETS = {
 "Settings": ["key", "value"],
 "Lines": ["line_id", "name", "source", "step_label", "steps_heading", "ask_question", "model_prompt", "model_help"],
 "Models": ["line_id", "model_id", "name", "note"],
 "Sections": ["line_id", "section_id", "label"],
 "Problems": ["line_id", "problem_id", "section_id", "name", "description", "images", "models"],
 "Steps": ["line_id", "problem_id", "order", "title", "remedy_keys", "caution_keys", "level", "note", "images", "hmi", "source", "internal_note"],
 "Remedies": ["line_id", "key", "text"],
 "Cautions": ["line_id", "key", "text"],
 "Glossary": ["line_id", "group", "term", "text", "image"],
 "Translations": ["english"],   # + one column per language id (filled at export)
 "Review": ["language", "term", "english", "note"],
}
README = """HOW TO EDIT THIS TROUBLESHOOTER
1. Edit the cream-shaded cells on any sheet. Add rows at the bottom of a table; delete rows to remove content.
2. Sheet guide:
   Settings  - app title, prompts, support contact (one key per row).
   Lines     - one row per machine line. step_label is what a step is called in the wizard ("Cause" or "Adjustment").
   Sections  - the groups shown in the left-hand picker, in the order listed.
   Problems  - one row per defect / symptom. 'images' = image ids from the images folder, with optional captions and model tags:  image11:Rough white marks | saffire_heat:Heat settings@saffire,ruby
               'models' = comma-separated model ids this problem applies to (blank = all models).
   Models    - one row per machine model within a line (the app asks the user to pick one when the line is chosen).
   Steps     - one row per cause/adjustment, in the order the wizard should try them. remedy_keys and caution_keys are comma-separated keys from the Remedies / Cautions sheets.
               level = operator (shown to everyone) or service (hidden behind a 'contact MSI Service' card in Operator mode).
               images = id:caption@model1,model2 | ... (omit @models to show for every model). hmi = model_id=screen path | model_id=... (use * for all). source = manual citation, or model_id=citation | ...
   Remedies  - the instruction text, one per key. A key can be reused by many steps.
   Cautions  - the warning text, one per key.
   Glossary  - optional picture glossary shown at the bottom of the Full matrix view.
   Translations - one row per English string, one column per language (de, es-ES, es-419 ...). Leave a cell blank to fall back to English.
               Languages offered in the app come from the Settings row 'languages' (e.g. en:English | de:Deutsch).
   Review    - terms to highlight in the app as 'under review' (language, term as it appears, English, note).
3. Pictures: drop JPG/PNG files into the images folder. The file name (without extension) is the image id you type in the sheets. Keep them under ~1500 px wide.
4. Rebuild:  python build_from_excel.py <this workbook> <images folder> <output folder>
   This writes <name>_artifact.html (for republishing the hosted page) and <name>_standalone.html (for SharePoint).
5. Alternatively send the workbook + any new pictures to Claude and ask for a rebuild and republish.
"""

def _list(s):  return [x.strip() for x in str(s or "").split(",") if x.strip()]
def _imgstr(im):
    s=f'{im["id"]}:{im.get("caption","")}'
    if im.get("models"): s+="@"+",".join(im["models"])
    return s
def _mapstr(v):
    if not v: return ""
    if isinstance(v,str): return v
    return " | ".join(f"{k}={x}" for k,x in v.items())
def _mapparse(s):
    s=str(s or "").strip()
    if not s: return None
    if "=" not in s: return s
    out={}
    for part in s.split(" | "):
        k,_,v=part.partition("="); out[k.strip()]=v.strip()
    return out
def _images(s):
    out=[]
    for part in str(s or "").split("|"):
        part=part.strip()
        if not part: continue
        iid,_,cap = part.partition(":")
        models=[]
        if "@" in cap: cap,_,m=cap.rpartition("@"); models=_list(m)
        d={"id":iid.strip(),"caption":cap.strip()}
        if models: d["models"]=models
        out.append(d)
    return out

def export(kb_path, xlsx_path, img_dir):
    kb=json.load(open(kb_path))
    wb=Workbook(); ws=wb.active; ws.title="README"
    ws["A1"]="README"; ws["A1"].font=Font(name=FONT,bold=True,size=14)
    for i,line in enumerate(README.strip().split("\n")):
        c=ws.cell(row=3+i,column=1,value=line); c.font=Font(name=FONT,size=10)
    ws.column_dimensions["A"].width=140
    rows={k:[] for k in SHEETS}
    for k,v in kb["ui"].items():
        rows["Settings"].append([k, " | ".join(v) if isinstance(v,list) else v])
    if kb.get("langs"): rows["Settings"].append(["languages"," | ".join(f'{l["id"]}:{l["label"]}' for l in kb["langs"])])
    tr=kb.get("tr") or {}; langs=[l["id"] for l in kb.get("langs",[]) if l["id"]!="en"]
    SHEETS["Translations"]=["english"]+langs
    keys=[]
    for lg in langs:
        for k in tr.get(lg,{}):
            if k not in keys: keys.append(k)
    for k in keys: rows["Translations"].append([k]+[tr.get(lg,{}).get(k,"") for lg in langs])
    for r in kb.get("review",[]): rows["Review"].append([r["lang"],r["term"],r["en"],r["note"]])
    for L in kb["lines"]:
        rows["Lines"].append([L["id"],L["name"],L.get("source",""),L.get("stepLabel",""),L.get("stepsHeading",""),L.get("askQuestion",""),L.get("modelPrompt",""),L.get("modelHelp","")])
        for m in L.get("models",[]): rows["Models"].append([L["id"],m["id"],m["name"],m.get("note","")])
        for sid,lab in (L.get("sections") or kb.get("sections") or {}).items():
            rows["Sections"].append([L["id"],sid,lab])
        for d in L["defects"]:
            rows["Problems"].append([L["id"],d["id"],d["section"],d["name"],d.get("desc",""),
                " | ".join(_imgstr(im) for im in d.get("images",[])), ", ".join(d.get("models",[]))])
            for i,c in enumerate(d["causes"],1):
                ims=c.get("images") or ([c["image"]] if c.get("image") else [])
                rows["Steps"].append([L["id"],d["id"],i,c["cause"],", ".join(c["remedies"]),", ".join(c["cautions"]),
                    c.get("level","operator"),c.get("note","")," | ".join(_imgstr(im) for im in ims),_mapstr(c.get("hmi")),_mapstr(c.get("source")),c.get("internal","")])
        for k,v in L["remedies"].items(): rows["Remedies"].append([L["id"],k,v])
        for k,v in L["cautions"].items(): rows["Cautions"].append([L["id"],k,v])
        for g in L.get("glossary",[]): rows["Glossary"].append([L["id"],g["group"],g["term"],g["text"],g.get("image","")])
    for name,cols in SHEETS.items():
        s=wb.create_sheet(name)
        for j,h in enumerate(cols,1):
            c=s.cell(row=1,column=j,value=h); c.font=Font(name=FONT,bold=True,color="FFFFFF"); c.fill=HEAD_FILL
        for i,r in enumerate(rows[name],2):
            for j,v in enumerate(r,1):
                c=s.cell(row=i,column=j,value=v); c.font=Font(name=FONT,size=10); c.fill=INPUT_FILL
                c.alignment=Alignment(wrap_text=True,vertical="top")
        widths={"text":80,"description":70,"value":60,"name":45,"title":45,"note":45,"source":40,"label":28,"images":40,"internal_note":40,"english":60,"de":60,"es-ES":60,"es-419":60,"term":24}
        for j,h in enumerate(cols,1): s.column_dimensions[get_column_letter(j)].width=widths.get(h,16)
        s.freeze_panes="A2"
    wb.save(xlsx_path)
    os.makedirs(img_dir,exist_ok=True)
    for iid,uri in (kb.get("images") or {}).items():
        head,b64=uri.split(",",1); ext="png" if "png" in head else "jpg"
        open(os.path.join(img_dir,f"{iid}.{ext}"),"wb").write(base64.b64decode(b64))
    print("wrote",xlsx_path,"and",len(kb.get("images") or {}),"images to",img_dir)

def import_(xlsx_path, img_dir, kb_path):
    wb=load_workbook(xlsx_path,data_only=True)
    def rows(name):
        s=wb[name]; hdr=[c.value for c in s[1]]
        for r in s.iter_rows(min_row=2,values_only=True):
            if not any(v not in (None,"") for v in r): continue
            yield dict(zip(hdr,r))
    ui={}; langs=[]
    for r in rows("Settings"):
        k,v=r["key"],r["value"]
        if k=="languages":
            for part in str(v).split("|"):
                i,_,lab=part.strip().partition(":"); langs.append({"id":i.strip(),"label":lab.strip() or i.strip()})
            continue
        ui[k]=[x.strip() for x in str(v).split("|")] if k=="introTips" else (v if v is not None else "")
    tr={}; review=[]
    if "Translations" in wb.sheetnames:
        hdr=[c.value for c in wb["Translations"][1]]
        for r in rows("Translations"):
            for lg in hdr[1:]:
                if lg and r.get(lg): tr.setdefault(lg,{})[str(r["english"])]=str(r[lg])
    if "Review" in wb.sheetnames:
        for r in rows("Review"): review.append({"lang":r["language"],"term":r["term"],"en":r["english"],"note":r["note"] or ""})
    lines={}
    for r in rows("Lines"):
        lines[r["line_id"]]={"id":r["line_id"],"name":r["name"],"source":r.get("source") or "","defects":[],"remedies":{},"cautions":{},"glossary":[],"sections":{}}
        for k,src in [("stepLabel","step_label"),("stepsHeading","steps_heading"),("askQuestion","ask_question"),("modelPrompt","model_prompt"),("modelHelp","model_help")]:
            if r.get(src): lines[r["line_id"]][k]=r[src]
    for r in rows("Sections"): lines[r["line_id"]]["sections"][r["section_id"]]=r["label"]
    if "Models" in wb.sheetnames:
        for r in rows("Models"): lines[r["line_id"]].setdefault("models",[]).append({"id":r["model_id"],"name":r["name"],"note":r.get("note") or ""})
    for r in rows("Remedies"): lines[r["line_id"]]["remedies"][str(r["key"])]=r["text"]
    for r in rows("Cautions"): lines[r["line_id"]]["cautions"][str(r["key"])]=r["text"]
    probs={}
    for r in rows("Problems"):
        d={"id":r["problem_id"],"section":r["section_id"],"name":r["name"],"desc":r.get("description") or "","images":_images(r.get("images")),"causes":[]}
        if r.get("models"): d["models"]=_list(r["models"])
        lines[r["line_id"]]["defects"].append(d); probs[(r["line_id"],r["problem_id"])]=d
    steps=sorted(rows("Steps"),key=lambda r:(r["line_id"],r["problem_id"],float(r["order"] or 0)))
    for r in steps:
        c={"cause":r["title"],"remedies":_list(r["remedy_keys"]),"cautions":_list(r["caution_keys"]),"level":(r.get("level") or "operator").strip().lower()}
        if r.get("note"): c["note"]=r["note"]
        if r.get("internal_note"): c["internal"]=r["internal_note"]
        if r.get("images"): c["images"]=_images(r["images"])
        if r.get("hmi"): c["hmi"]=_mapparse(r["hmi"])
        if r.get("source"): c["source"]=_mapparse(r["source"])
        L=lines[r["line_id"]]
        for k in c["remedies"]: assert k in L["remedies"], f"Steps: unknown remedy key {k} ({r['line_id']}/{r['problem_id']})"
        for k in c["cautions"]: assert k in L["cautions"], f"Steps: unknown caution key {k} ({r['line_id']}/{r['problem_id']})"
        probs[(r["line_id"],r["problem_id"])]["causes"].append(c)
    for r in rows("Glossary"):
        g={"group":r["group"],"term":r["term"],"text":r["text"]}
        if r.get("image"): g["image"]=str(r["image"]).strip()
        lines[r["line_id"]]["glossary"].append(g)
    # images from folder (resized + base64)
    images={}
    if os.path.isdir(img_dir):
        from PIL import Image; import io
        for f in sorted(os.listdir(img_dir)):
            iid,ext=os.path.splitext(f)
            if ext.lower() not in (".jpg",".jpeg",".png"): continue
            im=Image.open(os.path.join(img_dir,f)).convert("RGB"); im.thumbnail((900,900))
            b=io.BytesIO(); im.save(b,"JPEG",quality=72,optimize=True)
            images[iid]="data:image/jpeg;base64,"+base64.b64encode(b.getvalue()).decode()
    kb={"ui":ui,"sections":next(iter(lines.values()))["sections"] if lines else {},"lines":list(lines.values()),"images":images}
    if langs: kb["langs"]=langs
    if tr: kb["tr"]=tr
    if review: kb["review"]=review
    json.dump(kb,open(kb_path,"w"))
    print("wrote",kb_path,":",[(l["id"],len(l["defects"])) for l in kb["lines"]],len(images),"images")

if __name__=="__main__":
    if sys.argv[1]=="export": export(*sys.argv[2:5])
    else: import_(*sys.argv[2:5])
